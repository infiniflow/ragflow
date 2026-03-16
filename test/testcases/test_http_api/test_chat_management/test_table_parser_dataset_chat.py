#
#  Copyright 2025 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#
import os
import re
import tempfile

import pytest
from utils import wait_for

from common import (
    chat_completions,
    create_chat_assistant,
    create_session_with_chat_assistant,
    delete_chat_assistants,
    list_documents,
    parse_documents,
    upload_documents,
)


@wait_for(200, 1, "Document parsing timeout")
def wait_for_parsing_completion(auth, dataset_id, document_id=None):
    """
    Wait for document parsing to complete.

    Args:
        auth: Authentication object
        dataset_id: Dataset ID
        document_id: Optional specific document ID to wait for

    Returns:
        bool: True if parsing is complete, False otherwise
    """
    res = list_documents(auth, dataset_id)
    docs = res["data"]["docs"]

    if document_id is None:
        # Wait for all documents to complete
        for doc in docs:
            status = doc.get("run", "UNKNOWN")
            if status != "DONE":
                # print(f"[DEBUG] Document {doc.get('name', 'unknown')} status: {status}, progress: {doc.get('progress', 0)}%, msg: {doc.get('progress_msg', '')}")
                return False
        return True
    else:
        # Wait for specific document
        for doc in docs:
            if doc["id"] == document_id:
                status = doc.get("run", "UNKNOWN")
                # print(f"[DEBUG] Document {doc.get('name', 'unknown')} status: {status}, progress: {doc.get('progress', 0)}%, msg: {doc.get('progress_msg', '')}")
                if status == "DONE":
                    return True
                elif status == "FAILED":
                    pytest.fail(f"Document parsing failed: {doc}")
                    return False
        return False


# Test data
TEST_EXCEL_DATA = [
    ["employee_id", "name", "department", "salary"],
    ["E001", "Alice Johnson", "Engineering", "95000"],
    ["E002", "Bob Smith", "Marketing", "65000"],
    ["E003", "Carol Williams", "Engineering", "88000"],
    ["E004", "David Brown", "Sales", "72000"],
    ["E005", "Eva Davis", "HR", "68000"],
    ["E006", "Frank Miller", "Engineering", "102000"],
]

TEST_EXCEL_DATA_2 = [
    ["product", "price", "category"],
    ["Laptop", "999", "Electronics"],
    ["Mouse", "29", "Electronics"],
    ["Desk", "299", "Furniture"],
    ["Chair", "199", "Furniture"],
    ["Monitor", "399", "Electronics"],
    ["Keyboard", "79", "Electronics"],
]

DEFAULT_CHAT_PROMPT = (
    "You are a helpful assistant that answers questions about table data using SQL queries.\n\n"
    "Here is the knowledge base:\n{knowledge}\n\n"
    "Use this information to answer questions."
)


@pytest.mark.usefixtures("add_table_parser_dataset")
class TestTableParserDatasetChat:
    """
    Test table parser dataset chat functionality with Infinity backend.

    Verifies that:
    1. Excel files are uploaded and parsed correctly into table parser datasets
    2. Chat assistants can query the parsed table data via SQL
    3. Different types of queries work
    """

    @pytest.fixture(autouse=True)
    def setup_chat_assistant(self, HttpApiAuth, add_table_parser_dataset, request):
        """
        Setup fixture that runs before each test method.
        Creates chat assistant once and reuses it across all test cases.
        """
        # Only setup once (first time)
        if not hasattr(self.__class__, "chat_id"):
            self.__class__.dataset_id = add_table_parser_dataset
            self.__class__.auth = HttpApiAuth

            # Upload and parse Excel files once for all tests
            self._upload_and_parse_excel(HttpApiAuth, add_table_parser_dataset)

            # Create a single chat assistant and session for all tests
            chat_id, session_id = self._create_chat_assistant_with_session(HttpApiAuth, add_table_parser_dataset)
            self.__class__.chat_id = chat_id
            self.__class__.session_id = session_id

            # Store the total number of parametrize cases
            mark = request.node.get_closest_marker("parametrize")
            if mark:
                # Get the number of test cases from parametrize
                param_values = mark.args[1]
                self.__class__._total_tests = len(param_values)
            else:
                self.__class__._total_tests = 1

        yield

        # Teardown: cleanup chat assistant after all tests
        # Use a class-level counter to track tests
        if not hasattr(self.__class__, "_test_counter"):
            self.__class__._test_counter = 0
        self.__class__._test_counter += 1

        # Cleanup after all parametrize tests complete
        if self.__class__._test_counter >= self.__class__._total_tests:
            self._teardown_chat_assistant()

    def _teardown_chat_assistant(self):
        """Teardown method to clean up chat assistant."""
        if hasattr(self.__class__, "chat_id") and self.__class__.chat_id:
            try:
                delete_chat_assistants(self.__class__.auth, {"ids": [self.__class__.chat_id]})
            except Exception as e:
                print(f"[Teardown] Warning: Failed to delete chat assistant: {e}")

    @pytest.mark.p1
    @pytest.mark.parametrize(
        "question, expected_answer_pattern",
        [
            ("show me column of product", r"\|product\|Source"),
            ("which product has price 79", r"Keyboard"),
            ("How many rows in the dataset?", r"rows|count\(\*\)"),
            ("Show me all employees in Engineering department", r"(Alice|Carol|Frank)"),
        ],
    )
    def test_table_parser_dataset_chat(self, question, expected_answer_pattern):
        """
        Test that table parser dataset chat works correctly.
        """
        # Use class-level attributes (set by setup fixture)
        answer = self._ask_question(
            self.__class__.auth,
            self.__class__.chat_id,
            self.__class__.session_id,
            question
        )

        # Verify answer matches expected pattern if provided
        if expected_answer_pattern:
            self._assert_answer_matches_pattern(answer, expected_answer_pattern)
        else:
            # Just verify we got a non-empty answer
            assert answer and len(answer) > 0, "Expected non-empty answer"

    @staticmethod
    def _upload_and_parse_excel(auth, dataset_id):
        """
        Upload 2 Excel files and wait for parsing to complete.

        Returns:
            list: The document IDs of the uploaded files

        Raises:
            AssertionError: If upload or parsing fails
        """
        excel_file_paths = []
        document_ids = []
        try:
            # Create 2 temporary Excel files
            excel_file_paths.append(TestTableParserDatasetChat._create_temp_excel_file(TEST_EXCEL_DATA))
            excel_file_paths.append(TestTableParserDatasetChat._create_temp_excel_file(TEST_EXCEL_DATA_2))

            # Upload documents
            res = upload_documents(auth, dataset_id, excel_file_paths)
            assert res["code"] == 0, f"Failed to upload documents: {res}"

            for doc in res["data"]:
                document_ids.append(doc["id"])

            # Start parsing for all documents
            parse_payload = {"document_ids": document_ids}
            res = parse_documents(auth, dataset_id, parse_payload)
            assert res["code"] == 0, f"Failed to start parsing: {res}"

            # Wait for parsing completion for all documents
            for doc_id in document_ids:
                wait_for_parsing_completion(auth, dataset_id, doc_id)

            return document_ids

        finally:
            # Clean up temporary files
            for excel_file_path in excel_file_paths:
                if excel_file_path:
                    os.unlink(excel_file_path)

    @staticmethod
    def _create_temp_excel_file(data):
        """
        Create a temporary Excel file with the given table test data.

        Args:
            data: List of lists containing the Excel data

        Returns:
            str: Path to the created temporary file
        """
        from openpyxl import Workbook

        f = tempfile.NamedTemporaryFile(mode="wb", suffix=".xlsx", delete=False)
        f.close()

        wb = Workbook()
        ws = wb.active

        # Write test data to the worksheet
        for row_idx, row_data in enumerate(data, start=1):
            for col_idx, value in enumerate(row_data, start=1):
                ws.cell(row=row_idx, column=col_idx, value=value)

        wb.save(f.name)
        return f.name

    @staticmethod
    def _create_chat_assistant_with_session(auth, dataset_id):
        """
        Create a chat assistant and session for testing.

        Returns:
            tuple: (chat_id, session_id)
        """
        import uuid

        chat_payload = {
            "name": f"test_table_parser_dataset_chat_{uuid.uuid4().hex[:8]}",
            "dataset_ids": [dataset_id],
            "prompt_config": {
                "system": DEFAULT_CHAT_PROMPT,
                "parameters": [
                    {
                        "key": "knowledge",
                        "optional": True,
                        "value": "Use the table data to answer questions with SQL queries.",
                    }
                ],
            },
        }

        res = create_chat_assistant(auth, chat_payload)
        assert res["code"] == 0, f"Failed to create chat assistant: {res}"
        chat_id = res["data"]["id"]

        res = create_session_with_chat_assistant(auth, chat_id, {"name": f"test_session_{uuid.uuid4().hex[:8]}"})
        assert res["code"] == 0, f"Failed to create session: {res}"
        session_id = res["data"]["id"]

        return chat_id, session_id

    def _ask_question(self, auth, chat_id, session_id, question):
        """
        Send a question to the chat assistant and return the answer.

        Returns:
            str: The assistant's answer
        """
        payload = {
            "question": question,
            "stream": False,
            "session_id": session_id,
        }

        res_json = chat_completions(auth, chat_id, payload)
        assert res_json["code"] == 0, f"Chat completion failed: {res_json}"

        return res_json["data"]["answer"]

    def _assert_answer_matches_pattern(self, answer, pattern):
        """
        Assert that the answer matches the expected pattern.

        Args:
            answer: The actual answer from the chat assistant
            pattern: Regular expression pattern to match
        """
        assert re.search(pattern, answer, re.IGNORECASE), (
            f"Answer does not match expected pattern '{pattern}'.\n"
            f"Answer: {answer}"
        )
