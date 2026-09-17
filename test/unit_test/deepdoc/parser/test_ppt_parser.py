#
#  Copyright 2025 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#

"""Unit tests for RAGFlowPptParser table extraction.

A slide table is turned into "header: value" pairs by pairing every row with the
first one. A table that only has that one row therefore has nothing to pair with,
and used to come back empty.
"""

import importlib.util
import io
import os
import sys

import pytest

pptx = pytest.importorskip("pptx")
from pptx import Presentation  # noqa: E402
from pptx.util import Inches  # noqa: E402

# Import by file path so deepdoc/parser/__init__.py, which pulls in heavy
# dependencies, is not needed for this test.
_PARSER_PATH = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "..",
    "..",
    "..",
    "..",
    "deepdoc",
    "parser",
    "ppt_parser.py",
)
_spec = importlib.util.spec_from_file_location("ppt_parser_under_test", os.path.normpath(_PARSER_PATH))
_module = importlib.util.module_from_spec(_spec)
sys.modules["ppt_parser_under_test"] = _module
_spec.loader.exec_module(_module)
RAGFlowPptParser = _module.RAGFlowPptParser


def _deck_with_table(rows: int, cols: int = 2) -> bytes:
    """A one-slide deck with a text box and a `rows` x `cols` table."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])

    textbox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(4), Inches(0.5))
    textbox.text_frame.text = "SLIDE TEXT"

    table = slide.shapes.add_table(rows, cols, Inches(0.5), Inches(1.0), Inches(6), Inches(0.8 * rows)).table
    for row in range(rows):
        for col in range(cols):
            table.cell(row, col).text = f"R{row}C{col}"

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def _parse(payload: bytes) -> str:
    return RAGFlowPptParser()(payload, 0, 100000)[0]


def test_single_row_table_is_not_dropped():
    text = _parse(_deck_with_table(rows=1))

    assert "R0C0" in text
    assert "R0C1" in text
    assert "SLIDE TEXT" in text


def test_two_row_table_still_pairs_the_header_with_the_value():
    text = _parse(_deck_with_table(rows=2))

    assert "R0C0: R1C0; R0C1: R1C1" in text


def test_three_row_table_pairs_every_row_with_the_header():
    text = _parse(_deck_with_table(rows=3))

    assert "R0C0: R1C0; R0C1: R1C1" in text
    assert "R0C0: R2C0; R0C1: R2C1" in text
