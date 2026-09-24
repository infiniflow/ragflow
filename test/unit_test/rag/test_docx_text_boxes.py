#
#  Copyright 2025 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#

"""The DOCX chunkers dropped every text box.

A text box keeps its own ``w:p`` elements in a ``w:txbxContent`` nested inside
the drawing of a run, so neither ``Paragraph.text`` nor ``Run.text`` reaches it.
Every DOCX chunker walks exactly those two, which means callouts, pull quotes,
sidebars and diagram labels were silently missing from the index.

The tests below build the three shapes Word actually writes — a DrawingML
shape, a legacy VML shape, and the ``mc:AlternateContent`` pair that holds both
copies of the same box — and pin the extracted text for the shared helper and
for each of the four chunkers that use it.

The last group covers the rest of what the naive chunker's walk over the body
and ``Paragraph.text`` missed: block and inline content controls, custom XML,
tracked insertions, simple fields and smart tags.
"""

import importlib.util
import sys
import types
from io import BytesIO
from pathlib import Path
from unittest import mock

import pytest

_REPO_ROOT = Path(__file__).resolve().parents[3]

NS = (
    'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
    'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
    'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
    'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
    'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
    'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
    'xmlns:v="urn:schemas-microsoft-com:vml"'
)


def _box_paragraphs(*texts):
    return "".join(f"<w:p><w:r><w:t>{t}</w:t></w:r></w:p>" for t in texts)


def _drawing_box(*texts, inner=""):
    """A DrawingML text box — what Word 2010+ writes."""
    return f"""<w:drawing>
      <wp:inline distT="0" distB="0" distL="0" distR="0">
        <wp:extent cx="1000000" cy="500000"/>
        <wp:docPr id="1" name="Text Box 1"/>
        <a:graphic>
          <a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">
            <wps:wsp><wps:txbx><w:txbxContent>{_box_paragraphs(*texts)}{inner}</w:txbxContent></wps:txbx></wps:wsp>
          </a:graphicData>
        </a:graphic>
      </wp:inline>
    </w:drawing>"""


def _box_table(*rows):
    """A table inside a text box; its text lives in `w:p` below `w:tbl`."""
    cells = "".join(f"<w:tc>{_box_paragraphs(r)}</w:tc>" for r in rows)
    return f"<w:tbl><w:tr>{cells}</w:tr></w:tbl>"


def _vml_box(*texts):
    """A legacy VML text box — what older Word versions write."""
    return f"""<w:pict>
      <v:shape id="s1" type="#_x0000_t202">
        <v:textbox><w:txbxContent>{_box_paragraphs(*texts)}</w:txbxContent></v:textbox>
      </v:shape>
    </w:pict>"""


def _run(*payloads):
    return f"<w:r {NS}>{''.join(payloads)}</w:r>"


def _alternate_content_run(*texts):
    """The pair Word writes for a single text box: the same text twice."""
    return f"""<w:r {NS}>
      <mc:AlternateContent>
        <mc:Choice Requires="wps">{_drawing_box(*texts)}</mc:Choice>
        <mc:Fallback>{_vml_box(*texts)}</mc:Fallback>
      </mc:AlternateContent>
    </w:r>"""


def _build_docx(builder):
    from docx import Document

    doc = Document()
    builder(doc)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _anchor(paragraph, run_xml):
    from docx.oxml import parse_xml

    paragraph._p.append(parse_xml(run_xml))


# --------------------------------------------------------------------------- #
# Module loading
# --------------------------------------------------------------------------- #
#
# deepdoc/parser/__init__.py pulls in the vision stack, and the app modules pull
# in the LLM and storage stacks. Load docx_parser.py from source, publish the
# real class through a stubbed ``deepdoc.parser`` package, then load each
# chunker from source as well.


class _Dummy:
    def __init__(self, *a, **k):
        pass


def _stub(name, **attrs):
    mod = types.ModuleType(name)
    for key, value in attrs.items():
        setattr(mod, key, value)
    sys.modules[name] = mod
    return mod


def _load(module_name, relative_path):
    spec = importlib.util.spec_from_file_location(module_name, _REPO_ROOT / relative_path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def docx_modules():
    saved = {name: sys.modules.get(name) for name in list(sys.modules)}
    try:
        for name in ("xgboost", "pdfplumber", "huggingface_hub", "pypdf", "sklearn", "sklearn.cluster", "sklearn.metrics", "deepdoc.vision"):
            sys.modules.setdefault(name, mock.MagicMock())

        docx_parser = _load("deepdoc.parser.docx_parser_under_test", "deepdoc/parser/docx_parser.py")
        real_parser = docx_parser.RAGFlowDocxParser

        _stub(
            "deepdoc.parser",
            DocxParser=real_parser,
            PdfParser=_Dummy,
            HtmlParser=_Dummy,
            ExcelParser=_Dummy,
            EpubParser=_Dummy,
            JsonParser=_Dummy,
            MarkdownElementExtractor=_Dummy,
            MarkdownParser=_Dummy,
            TxtParser=_Dummy,
        )
        _stub("deepdoc.parser.utils", get_text=lambda *a, **k: "", extract_pdf_outlines=lambda *a, **k: [])
        _stub(
            "deepdoc.parser.figure_parser",
            VisionFigureParser=_Dummy,
            vision_figure_parser_pdf_wrapper=lambda *a, **k: [],
            vision_figure_parser_docx_wrapper=lambda *a, **k: [],
            vision_figure_parser_docx_wrapper_naive=lambda *a, **k: [],
        )
        _stub("deepdoc.parser.pdf_parser", PlainParser=_Dummy, VisionParser=_Dummy, RAGFlowPdfParser=_Dummy)
        _stub("deepdoc.parser.docling_parser", DoclingParser=_Dummy)
        _stub("deepdoc.parser.monkeyocrv2_parser", MonkeyOCRv2Parser=_Dummy)
        _stub("deepdoc.parser.tcadp_parser", TCADPParser=_Dummy)
        _stub(
            "api.db.joint_services.tenant_model_service",
            get_composite_model_name_by_id=lambda x: x,
            get_first_provider_model_name=lambda *a, **k: None,
            resolve_model_config=lambda *a, **k: None,
            get_tenant_default_model_by_type=lambda *a, **k: None,
            ensure_mineru_from_env=lambda *a, **k: None,
            ensure_opendataloader_from_env=lambda *a, **k: None,
            ensure_paddleocr_from_env=lambda *a, **k: None,
            ensure_somark_from_env=lambda *a, **k: None,
        )
        _stub("api.db.services.llm_service", LLMBundle=_Dummy)
        _stub(
            "rag.utils.file_utils",
            extract_embed_file=lambda *a, **k: None,
            extract_links_from_pdf=lambda *a, **k: [],
            extract_links_from_docx=lambda *a, **k: [],
            extract_html=lambda *a, **k: (None, None),
        )
        _stub("common.float_utils", get_float=lambda *a, **k: 0.0, normalize_overlapped_percent=lambda x: x)
        _stub("common.text_utils", normalize_arabic_presentation_forms=lambda x: x)
        _stub("rag.app.naive", by_plaintext=lambda *a, **k: ([], [], None), PARSERS={})
        _stub(
            "common.parser_config_utils",
            normalize_layout_recognizer=lambda x: (x, None),
            has_mineru_options=lambda *a, **k: False,
            is_tenant_model_id=lambda *a, **k: False,
            MINERU_OPTION_KEYS=(),
        )

        modules = types.SimpleNamespace(
            parser=real_parser,
            laws=_load("rag.app.laws_under_test", "rag/app/laws.py").Docx,
            qa=_load("rag.app.qa_under_test", "rag/app/qa.py").Docx,
            manual=_load("rag.app.manual_under_test", "rag/app/manual.py").Docx,
            naive=_load("rag.app.naive_under_test", "rag/app/naive.py").Docx,
        )
        yield modules
    finally:
        for name in list(sys.modules):
            if name not in saved:
                del sys.modules[name]
        sys.modules.update(saved)


# --------------------------------------------------------------------------- #
# deepdoc.parser.docx_parser.RAGFlowDocxParser
# --------------------------------------------------------------------------- #


@pytest.mark.p2
def test_drawing_text_box_is_extracted(docx_modules):
    def builder(d):
        _anchor(d.add_paragraph("PARAGRAPH TEXT"), _run(_drawing_box("CALLOUT LINE ONE", "CALLOUT LINE TWO")))
        d.add_paragraph("AFTER")

    secs, _ = docx_modules.parser()(_build_docx(builder))
    texts = [text for text, _style in secs]

    assert texts == ["PARAGRAPH TEXT", "CALLOUT LINE ONE\nCALLOUT LINE TWO", "AFTER"]


@pytest.mark.p2
def test_text_box_is_not_glued_to_its_paragraph(docx_modules):
    """Sweeping .//w:t over the paragraph reaches the box but splices it into
    the sentence; the box has to stay a block of its own."""

    def builder(d):
        _anchor(d.add_paragraph("PARAGRAPH TEXT"), _run(_drawing_box("CALLOUT")))

    secs, _ = docx_modules.parser()(_build_docx(builder))

    assert not any("PARAGRAPH TEXTCALLOUT" in text for text, _style in secs)


@pytest.mark.p2
def test_vml_text_box_is_extracted(docx_modules):
    def builder(d):
        _anchor(d.add_paragraph("PARAGRAPH TEXT"), _run(_vml_box("LEGACY CALLOUT")))

    secs, _ = docx_modules.parser()(_build_docx(builder))

    assert "LEGACY CALLOUT" in [text for text, _style in secs]


@pytest.mark.p2
def test_alternate_content_text_box_is_read_once(docx_modules):
    """Word stores the same box twice — a DrawingML shape under mc:Choice and a
    VML shape under mc:Fallback. Reading both duplicates the text."""

    def builder(d):
        _anchor(d.add_paragraph("PARAGRAPH TEXT"), _alternate_content_run("CALLOUT"))

    secs, _ = docx_modules.parser()(_build_docx(builder))

    assert [text for text, _style in secs].count("CALLOUT") == 1


@pytest.mark.p2
def test_nested_text_box_is_read_once(docx_modules):
    def builder(d):
        _anchor(d.add_paragraph("PARAGRAPH TEXT"), _run(_drawing_box("OUTER", inner=_vml_box("INNER"))))

    secs, _ = docx_modules.parser()(_build_docx(builder))
    texts = [text for text, _style in secs]

    assert texts.count("OUTER") == 1
    assert texts.count("INNER") == 1


@pytest.mark.p2
def test_table_inside_a_text_box_is_extracted(docx_modules):
    """A text box can hold a table; its cell text sits below `w:tbl`, so only
    walking the box's direct `w:p` children would miss it."""

    def builder(d):
        _anchor(d.add_paragraph("PARAGRAPH TEXT"), _run(_drawing_box("BOX TITLE", inner=_box_table("CELL A", "CELL B"))))

    secs, _ = docx_modules.parser()(_build_docx(builder))
    texts = [text for text, _style in secs]

    assert texts == ["PARAGRAPH TEXT", "BOX TITLE\nCELL A\nCELL B"]


@pytest.mark.p2
def test_document_without_text_boxes_is_unchanged(docx_modules):
    """Guard: the extraction must not touch a document that has no text box."""

    def builder(d):
        d.add_paragraph("FIRST")
        d.add_paragraph("SECOND")

    secs, _ = docx_modules.parser()(_build_docx(builder))

    assert [text for text, _style in secs] == ["FIRST", "SECOND"]


@pytest.mark.p2
def test_extract_text_boxes_tolerates_a_paragraph_double(docx_modules):
    """Existing tests hand the parsers fakes whose ``_element`` is a plain list."""
    fake = types.SimpleNamespace(_element=[])

    assert docx_modules.parser.extract_text_boxes(fake) == []
    assert docx_modules.parser.extract_text_boxes(types.SimpleNamespace()) == []


@pytest.mark.p2
def test_text_box_outside_the_page_range_is_skipped(docx_modules):
    def builder(d):
        _anchor(d.add_paragraph("PARAGRAPH TEXT"), _run(_drawing_box("CALLOUT")))

    secs, _ = docx_modules.parser()(_build_docx(builder), from_page=1, to_page=2)

    assert "CALLOUT" not in [text for text, _style in secs]


def _break_then_box_paragraph(d):
    """A paragraph whose first run ends page 0 and whose second run anchors a box."""
    paragraph = d.add_paragraph()
    _anchor(paragraph, f"<w:r {NS}><w:t>PAGE ZERO TEXT</w:t><w:lastRenderedPageBreak/></w:r>")
    _anchor(paragraph, _run(_drawing_box("BOX ON PAGE ONE")))


@pytest.mark.p2
def test_text_box_is_placed_on_the_page_of_its_anchor(docx_modules):
    """A paragraph can span pages; the box belongs to the page of the run that
    anchors it, not to the page the paragraph started on."""
    secs, _ = docx_modules.parser()(_build_docx(_break_then_box_paragraph), from_page=1, to_page=2)

    assert "BOX ON PAGE ONE" in [text for text, _style in secs]


@pytest.mark.p2
def test_text_box_on_a_later_page_is_left_out_of_an_earlier_range(docx_modules):
    secs, _ = docx_modules.parser()(_build_docx(_break_then_box_paragraph), from_page=0, to_page=1)
    texts = [text for text, _style in secs]

    assert "PAGE ZERO TEXT" in texts
    assert "BOX ON PAGE ONE" not in texts


@pytest.mark.p2
def test_text_box_found_only_under_a_fallback_is_kept(docx_modules):
    """`mc:Fallback` is a general container. Only the copy of a box that its
    `mc:Choice` already holds is a duplicate; a box found only there is the one
    copy the document has."""
    run = f"""<w:r {NS}>
      <mc:AlternateContent>
        <mc:Choice Requires="wps"><w:t>no text box here</w:t></mc:Choice>
        <mc:Fallback>{_vml_box("ONLY IN THE FALLBACK")}</mc:Fallback>
      </mc:AlternateContent>
    </w:r>"""

    def builder(d):
        _anchor(d.add_paragraph("PARAGRAPH TEXT"), run)

    secs, _ = docx_modules.parser()(_build_docx(builder))

    assert "ONLY IN THE FALLBACK" in [text for text, _style in secs]


# --------------------------------------------------------------------------- #
# The chunkers
# --------------------------------------------------------------------------- #


@pytest.mark.p2
def test_laws_docx_keeps_text_box_in_its_section(docx_modules):
    def builder(d):
        d.add_heading("Chapter 1 General Provisions", level=1)
        d.add_heading("Article 2 Fee Schedule", level=2)
        _anchor(d.add_paragraph("The applicable fees are as follows:"), _run(_drawing_box("NOTE: fees are waived for students")))

    chunks = docx_modules.laws()("law.docx", _build_docx(builder))

    assert any("fees are waived for students" in c for c in chunks)
    # A text box carries no heading level, so it stays leaf content of the
    # enclosing section and keeps that section's title path as context.
    callout_chunk = next(c for c in chunks if "fees are waived for students" in c)
    assert "Article 2 Fee Schedule" in callout_chunk


@pytest.mark.p2
def test_laws_docx_keeps_text_box_anchored_in_an_empty_paragraph(docx_modules):
    def builder(d):
        d.add_heading("Chapter 1", level=1)
        _anchor(d.add_paragraph(""), _run(_drawing_box("FLOATING CALLOUT")))

    chunks = docx_modules.laws()("law.docx", _build_docx(builder))

    assert any("FLOATING CALLOUT" in c for c in chunks)


@pytest.mark.p2
def test_naive_docx_keeps_text_box_as_its_own_block(docx_modules):
    def builder(d):
        _anchor(d.add_paragraph("Connect the power cable."), _run(_drawing_box("WARNING: disconnect the battery first")))
        d.add_paragraph("AFTER")

    lines = docx_modules.naive()("manual.docx", _build_docx(builder))

    assert [text for text, _img, _tbl in lines] == ["Connect the power cable.", "WARNING: disconnect the battery first", "AFTER"]


@pytest.mark.p2
def test_qa_docx_keeps_text_box_in_the_answer(docx_modules):
    def builder(d):
        d.add_heading("How do I reset my password?", level=1)
        _anchor(d.add_paragraph("Open the settings page."), _run(_drawing_box("TIP: the link expires after one hour")))

    qai_list, _tbls = docx_modules.qa()("qa.docx", _build_docx(builder))

    assert any("link expires after one hour" in answer for _q, answer, _img in qai_list)


@pytest.mark.p2
def test_manual_docx_keeps_text_box_in_the_section(docx_modules):
    def builder(d):
        d.add_heading("1 Installation", level=1)
        _anchor(d.add_paragraph("Connect the power cable."), _run(_drawing_box("WARNING: disconnect the battery first")))

    ti_list, _tbls = docx_modules.manual()("manual.docx", _build_docx(builder))

    assert any("disconnect the battery first" in text for text, _img in ti_list)


# --------------------------------------------------------------------------- #
# Content controls, tracked changes and other runs below the paragraph
# --------------------------------------------------------------------------- #

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
REVISION = 'w:author="a" w:date="2024-01-01T00:00:00Z"'


def _append(document, xml):
    from docx.oxml import parse_xml

    body = document.element.body
    body.insert(len(body) - 1, parse_xml(xml))  # ahead of the final sectPr


def _paragraph(inner):
    return f"<w:p {W}>{inner}</w:p>"


def _text(text):
    return f'<w:r><w:t xml:space="preserve">{text}</w:t></w:r>'


def _naive(docx_modules, builder):
    return docx_modules.naive()("doc.docx", _build_docx(builder))


@pytest.mark.p2
def test_block_content_controls_are_read(docx_modules):
    def builder(d):
        d.add_paragraph("BEFORE")
        _append(d, f'<w:sdt {W}><w:sdtPr/><w:sdtContent><w:p>{_text("IN A CONTROL")}</w:p></w:sdtContent></w:sdt>')
        _append(d, f'<w:customXml {W} w:element="clause"><w:p>{_text("IN CUSTOM XML")}</w:p></w:customXml>')
        d.add_paragraph("AFTER")

    assert [text for text, _image, _table in _naive(docx_modules, builder)] == ["BEFORE", "IN A CONTROL", "IN CUSTOM XML", "AFTER"]


@pytest.mark.p2
def test_runs_nested_in_a_paragraph_are_read(docx_modules):
    def builder(d):
        _append(d, _paragraph(_text("Signed by ") + f"<w:sdt><w:sdtPr/><w:sdtContent>{_text('Jane Doe')}</w:sdtContent></w:sdt>"))
        _append(
            d,
            _paragraph(
                _text("Payment is due in ")
                # A deleted tab is not delText, so it would still leak into the text.
                + f'<w:del w:id="1" {REVISION}><w:r><w:tab/><w:delText>60</w:delText></w:r></w:del>'
                + f'<w:ins w:id="2" {REVISION}>{_text("30")}</w:ins>'
                + _text(" days")
            ),
        )
        _append(d, _paragraph(_text("Updated ") + f'<w:fldSimple w:instr="DATE">{_text("2024-09-30")}</w:fldSimple>'))
        _append(d, _paragraph(f'<w:smartTag w:uri="u" w:element="place">{_text("Reutlingen")}</w:smartTag>'))

    assert [text for text, _image, _table in _naive(docx_modules, builder)] == [
        "Signed by Jane Doe",
        "Payment is due in 30 days",
        "Updated 2024-09-30",
        "Reutlingen",
    ]


@pytest.mark.p2
def test_runs_that_are_not_the_text_stay_out(docx_modules):
    """A move's old place, the pronunciation guide of ruby text, and the
    fallback copy of alternate content are read neither in place of nor in
    addition to the text they shadow."""
    mc = 'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'

    def builder(d):
        _append(d, _paragraph(f'<w:moveFrom w:id="1" {REVISION}>{_text("OLD PLACE")}</w:moveFrom><w:moveTo w:id="2" {REVISION}>{_text("NEW PLACE")}</w:moveTo>'))
        _append(d, _paragraph(f'<w:r><w:ruby><w:rubyPr/><w:rt>{_text("kan")}</w:rt><w:rubyBase>{_text("漢")}</w:rubyBase></w:ruby></w:r>'))
        alternate = f'<mc:AlternateContent><mc:Choice Requires="w14">{_text("CHOICE")}</mc:Choice><mc:Fallback>{_text("FALLBACK")}</mc:Fallback></mc:AlternateContent>'
        _append(d, f"<w:p {W} {mc}>{alternate}</w:p>")

    assert [text for text, _image, _table in _naive(docx_modules, builder)] == ["NEW PLACE", "漢", "CHOICE"]


@pytest.mark.p2
def test_a_table_in_a_control_is_read_and_later_tables_keep_their_titles(docx_modules):
    """The table title lookup counts tables the same way the walk does."""

    def builder(d):
        d.add_heading("Alpha", level=1)
        table = f"<w:tbl><w:tblPr/><w:tblGrid><w:gridCol/></w:tblGrid><w:tr><w:tc><w:p>{_text('T1')}</w:p></w:tc></w:tr></w:tbl>"
        _append(d, f"<w:sdt {W}><w:sdtPr/><w:sdtContent>{table}</w:sdtContent></w:sdt>")
        d.add_heading("Beta", level=1)
        d.add_table(rows=1, cols=1).cell(0, 0).text = "T2"

    assert [table for _text_, _image, table in _naive(docx_modules, builder) if table] == [
        "<table><caption>Table Location: doc > Alpha</caption><tr><td>T1</td></tr></table>",
        "<table><caption>Table Location: doc > Beta</caption><tr><td>T2</td></tr></table>",
    ]


@pytest.mark.p2
def test_a_table_caption_reads_a_heading_held_in_a_content_control(docx_modules):
    """The caption is built from the same heading text the chunk shows."""

    def builder(d):
        from docx.oxml import parse_xml

        heading = d.add_heading("", level=1)
        heading._p.append(parse_xml(f"<w:sdt {W}><w:sdtPr/><w:sdtContent>{_text('Results')}</w:sdtContent></w:sdt>"))
        d.add_table(rows=1, cols=1).cell(0, 0).text = "T1"

    assert [table for _text_, _image, table in _naive(docx_modules, builder) if table] == [
        "<table><caption>Table Location: doc > Results</caption><tr><td>T1</td></tr></table>",
    ]


@pytest.mark.p2
@pytest.mark.parametrize(
    ("page_break", "first_page"),
    [
        (f'<w:ins w:id="9" {REVISION}><w:r><w:br w:type="page"/></w:r></w:ins>', "First page"),
        ("<w:sdt><w:sdtPr/><w:sdtContent><w:r><w:lastRenderedPageBreak/><w:t>.</w:t></w:r></w:sdtContent></w:sdt>", "First page."),
    ],
    ids=["page-break-in-an-insertion", "rendered-break-in-a-control"],
)
def test_page_breaks_in_nested_runs_are_counted(docx_modules, page_break, first_page):
    """A page break sits in a run the text is read from, so it moves what follows to the next page."""

    def builder(d):
        _append(d, _paragraph(_text("First page") + page_break))
        d.add_paragraph("Second page")

    lines = docx_modules.naive()("doc.docx", _build_docx(builder), from_page=0, to_page=1)

    assert [text for text, _image, _table in lines] == [first_page]


@pytest.mark.p2
def test_a_plain_paragraph_reads_exactly_as_before(docx_modules):
    """Direct runs, a tab, a line break and a hyperlink: the same text as Paragraph.text."""

    def builder(d):
        paragraph = d.add_paragraph("one\ttwo")
        paragraph.add_run().add_break()
        paragraph.add_run("three")
        _append(d, _paragraph(_text("see ") + f'<w:hyperlink w:anchor="x">{_text("the link")}</w:hyperlink>'))

    from docx import Document

    paragraphs = Document(BytesIO(_build_docx(builder))).paragraphs
    assert [docx_modules.parser.paragraph_text(p) for p in paragraphs] == [p.text for p in paragraphs]
    assert [p.text for p in paragraphs] == ["one\ttwo\nthree", "see the link"]
