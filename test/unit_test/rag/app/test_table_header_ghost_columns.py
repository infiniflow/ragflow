#
#  Copyright 2025 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#

"""Header width must follow real data, not the worksheet's declared dimension.

Excel records a style applied to a far column in the sheet ``<dimension>`` tag
even when no value is written there. openpyxl materialises a cell object for
every column that tag declares, so a four-column sheet can produce 16384 cells
per row. Header parsing used to name every one of those cells, turning the
dataset's field_map/table_column_names into thousands of ``Column_N`` entries.

The header helpers are extracted from the source with ``ast`` so the test does
not pull in rag.app.table's tokenizer/LLM import chain.
"""

from __future__ import annotations

import ast
import logging
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")
from openpyxl import Workbook
from openpyxl.styles import PatternFill

pytestmark = pytest.mark.p2

GHOST_COLUMN = 16384


def _load_excel_class():
    src_path = Path(__file__).resolve().parents[4] / "rag" / "app" / "table.py"
    tree = ast.parse(src_path.read_text(encoding="utf-8"))

    cls_node = next(node for node in tree.body if isinstance(node, ast.ClassDef) and node.name == "Excel")
    wanted = {
        "_used_column_count",
        "_parse_simple_headers",
        "_build_hierarchical_headers",
        "_get_merged_cell_value",
        "_is_valid_header_part",
    }
    nodes = [n for n in cls_node.body if isinstance(n, ast.FunctionDef) and n.name in wanted]
    assert {n.name for n in nodes} == wanted, f"missing methods: {wanted - {n.name for n in nodes}}"

    module = ast.Module(body=nodes, type_ignores=[])
    ns = {}
    exec(compile(module, str(src_path), "exec"), ns)  # noqa: S102 - exec'ing AST-extracted methods from our own source
    return type("Excel", (), ns)


Excel = _load_excel_class()


def _load_rows_limited():
    """Extract RAGFlowExcelParser._get_rows_limited and its helpers."""
    src_path = Path(__file__).resolve().parents[4] / "deepdoc" / "parser" / "excel_parser.py"
    tree = ast.parse(src_path.read_text(encoding="utf-8"))

    cls_node = next(node for node in tree.body if isinstance(node, ast.ClassDef) and node.name == "RAGFlowExcelParser")
    wanted = {"_get_actual_row_count", "_get_actual_column_count", "_get_rows_limited"}
    nodes = [n for n in cls_node.body if isinstance(n, ast.FunctionDef) and n.name in wanted]
    assert {n.name for n in nodes} == wanted, f"missing methods: {wanted - {n.name for n in nodes}}"

    ns = {"logging": logging}
    exec(compile(ast.Module(body=nodes, type_ignores=[]), str(src_path), "exec"), ns)  # noqa: S102 - own source
    cls = type("RAGFlowExcelParser", (), ns)
    ns["RAGFlowExcelParser"] = cls
    return cls


RAGFlowExcelParser = _load_rows_limited()


def _workbook_with_ghost_style(rows_of_data=None):
    """Build a 4-column sheet whose dimension is inflated by a far-column style."""
    wb = Workbook()
    ws = wb.active
    for row in rows_of_data or [
        ["ID", "Name", "Amount", "Status"],
        ["A-1", "alice", "10", "paid"],
        ["A-2", "bob", "20", "pending"],
    ]:
        ws.append(row)
    ws.cell(row=ws.max_row, column=GHOST_COLUMN).fill = PatternFill(fill_type="solid", fgColor="FFFF00")
    return ws


def test_rows_limited_bounds_width_to_used_columns():
    """The read layer must not materialise declared-but-empty columns."""
    ws = _workbook_with_ghost_style()

    rows = RAGFlowExcelParser._get_rows_limited(ws)

    assert len(rows) == 3
    assert all(len(row) == 4 for row in rows)


def test_rows_limited_keeps_trailing_data_column_with_blank_header():
    ws = Workbook().active
    ws["A1"] = "ID"
    ws["B1"] = "Name"
    ws["A2"] = "A-1"
    ws["B2"] = "alice"
    ws["C2"] = "orphan"

    rows = RAGFlowExcelParser._get_rows_limited(ws)

    assert len(rows[0]) == 3


def test_rows_limited_returns_empty_for_blank_sheet():
    assert RAGFlowExcelParser._get_rows_limited(Workbook().active) == []


def test_ghost_columns_are_not_named():
    ws = _workbook_with_ghost_style()
    rows = list(ws.iter_rows())

    # Precondition: openpyxl really does materialise the inflated width.
    assert len(rows[0]) == GHOST_COLUMN

    headers, header_rows = Excel()._parse_simple_headers(rows)

    assert header_rows == 1
    assert headers == ["ID", "Name", "Amount", "Status"]


def test_blank_cells_inside_used_range_keep_placeholder_names():
    """A gap between real columns must not shift later columns left (PR #9225)."""
    ws = Workbook().active
    ws["A1"] = "ID"
    ws["C1"] = "Status"
    ws["A2"] = "A-1"
    ws["C2"] = "paid"

    headers, _ = Excel()._parse_simple_headers(list(ws.iter_rows()))

    assert headers == ["ID", "Column_2", "Status"]


def test_trailing_column_with_data_below_blank_header_is_kept():
    """Bounding by the header row alone would drop a real column."""
    ws = Workbook().active
    ws["A1"] = "ID"
    ws["B1"] = "Name"
    ws["A2"] = "A-1"
    ws["B2"] = "alice"
    ws["C2"] = "orphan"  # value present, header cell empty

    headers, _ = Excel()._parse_simple_headers(list(ws.iter_rows()))

    assert headers == ["ID", "Name", "Column_3"]


def test_hierarchical_headers_are_bounded_by_used_range():
    ws = _workbook_with_ghost_style()
    rows = list(ws.iter_rows())

    headers = Excel()._build_hierarchical_headers(ws, rows, 1)

    assert headers == ["ID", "Name", "Amount", "Status"]


def test_used_column_count_ignores_whitespace_only_cells():
    ws = Workbook().active
    ws["A1"] = "ID"
    ws["B1"] = "   "
    ws["A2"] = "A-1"

    assert Excel._used_column_count(list(ws.iter_rows())) == 1
