#
#  Copyright 2025 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#

"""Three DOCX chunkers built their table HTML without escaping the cells.

`rag/app/laws.py` escapes them (`test_laws_docx_tables.py` pins it), and so does
`deepdoc/parser/excel_parser.py`. `naive.py`, `manual.py` and `qa.py` did not, so
a cell holding `<`, `>` or `&` -- a comparison, a formula, an XML snippet, a
company name -- was written into the table as markup.
"""

import importlib.util
import sys
import types
from io import BytesIO
from pathlib import Path
from unittest import mock

import pytest

_REPO_ROOT = Path(__file__).resolve().parents[3]

CELL_A = "a < b & c > d"
CELL_B = "<img src=x onerror=alert(1)>"
ESCAPED_A = "a &lt; b &amp; c &gt; d"
ESCAPED_B = "&lt;img src=x onerror=alert(1)&gt;"


class _Dummy:
    def __init__(self, *a, **k):
        pass


def _stub(name, **attrs):
    mod = types.ModuleType(name)
    for key, value in attrs.items():
        setattr(mod, key, value)
    sys.modules[name] = mod
    return mod


def _load(module_name, relative_path):
    spec = importlib.util.spec_from_file_location(module_name, _REPO_ROOT / relative_path)
    module = importlib.util.module_from_spec(spec)
    sys.modules[module_name] = module
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def chunkers():
    """The DOCX chunkers, with the vision / LLM / storage siblings stubbed out."""
    saved = {name: sys.modules.get(name) for name in list(sys.modules)}
    try:
        for name in ("xgboost", "pdfplumber", "huggingface_hub", "pypdf", "sklearn", "sklearn.cluster", "sklearn.metrics", "deepdoc.vision"):
            sys.modules.setdefault(name, mock.MagicMock())

        docx_parser = _load("deepdoc.parser.docx_parser_under_test", "deepdoc/parser/docx_parser.py")

        _stub(
            "deepdoc.parser",
            DocxParser=docx_parser.RAGFlowDocxParser,
            PdfParser=_Dummy,
            HtmlParser=_Dummy,
            ExcelParser=_Dummy,
            EpubParser=_Dummy,
            JsonParser=_Dummy,
            MarkdownElementExtractor=_Dummy,
            MarkdownParser=_Dummy,
            TxtParser=_Dummy,
        )
        _stub("deepdoc.parser.utils", get_text=lambda *a, **k: "", extract_pdf_outlines=lambda *a, **k: [])
        _stub(
            "deepdoc.parser.figure_parser",
            VisionFigureParser=_Dummy,
            vision_figure_parser_pdf_wrapper=lambda *a, **k: [],
            vision_figure_parser_docx_wrapper=lambda *a, **k: [],
            vision_figure_parser_docx_wrapper_naive=lambda *a, **k: [],
        )
        _stub("deepdoc.parser.pdf_parser", PlainParser=_Dummy, VisionParser=_Dummy, RAGFlowPdfParser=_Dummy)
        _stub("deepdoc.parser.docling_parser", DoclingParser=_Dummy)
        _stub("deepdoc.parser.monkeyocrv2_parser", MonkeyOCRv2Parser=_Dummy)
        _stub("deepdoc.parser.tcadp_parser", TCADPParser=_Dummy)
        _stub(
            "api.db.joint_services.tenant_model_service",
            get_composite_model_name_by_id=lambda x: x,
            get_first_provider_model_name=lambda *a, **k: None,
            resolve_model_config=lambda *a, **k: None,
            get_tenant_default_model_by_type=lambda *a, **k: None,
            ensure_mineru_from_env=lambda *a, **k: None,
            ensure_opendataloader_from_env=lambda *a, **k: None,
            ensure_paddleocr_from_env=lambda *a, **k: None,
            ensure_somark_from_env=lambda *a, **k: None,
        )
        _stub("api.db.services.llm_service", LLMBundle=_Dummy)
        _stub(
            "rag.utils.file_utils",
            extract_embed_file=lambda *a, **k: None,
            extract_links_from_pdf=lambda *a, **k: [],
            extract_links_from_docx=lambda *a, **k: [],
            extract_html=lambda *a, **k: (None, None),
        )
        _stub("common.float_utils", get_float=lambda *a, **k: 0.0, normalize_overlapped_percent=lambda x: x)
        _stub("common.text_utils", normalize_arabic_presentation_forms=lambda x: x)
        _stub("rag.app.naive", by_plaintext=lambda *a, **k: ([], [], None), PARSERS={})
        _stub(
            "common.parser_config_utils",
            normalize_layout_recognizer=lambda x: (x, None),
            has_mineru_options=lambda *a, **k: False,
            is_tenant_model_id=lambda *a, **k: False,
            MINERU_OPTION_KEYS=(),
        )

        yield types.SimpleNamespace(
            laws=_load("rag.app.laws_under_test", "rag/app/laws.py").Docx,
            qa=_load("rag.app.qa_under_test", "rag/app/qa.py").Docx,
            manual=_load("rag.app.manual_under_test", "rag/app/manual.py").Docx,
            naive=_load("rag.app.naive_under_test", "rag/app/naive.py").Docx,
        )
    finally:
        for name in list(sys.modules):
            if name not in saved:
                del sys.modules[name]
        sys.modules.update(saved)


def _docx_with_table(heading="Heading"):
    from docx import Document

    doc = Document()
    doc.add_heading(heading, level=1)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = CELL_A
    table.cell(0, 1).text = CELL_B

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _assert_escaped(html):
    assert ESCAPED_A in html
    assert ESCAPED_B in html
    # The markup a cell carried must not reach the table as markup.
    assert "<img" not in html
    assert "<td>a < b" not in html


@pytest.mark.p2
def test_naive_docx_escapes_cell_html(chunkers):
    lines = chunkers.naive()("table.docx", _docx_with_table())

    tables = [table for _text, _image, table in lines if table]
    assert tables
    _assert_escaped(tables[0])


@pytest.mark.p2
def test_naive_docx_escapes_the_table_caption(chunkers):
    """The caption carries the enclosing heading, which is document text too."""
    lines = chunkers.naive()("table.docx", _docx_with_table(heading="Q1 < Q2 & Q3"))

    tables = [table for _text, _image, table in lines if table]
    assert tables
    assert "<caption>" in tables[0]
    assert "Q1 &lt; Q2 &amp; Q3" in tables[0]
    assert "Q1 < Q2" not in tables[0]


@pytest.mark.p2
def test_manual_docx_escapes_cell_html(chunkers):
    _ti_list, tbls = chunkers.manual()("table.docx", _docx_with_table())

    assert tbls
    _assert_escaped(tbls[0][0][1])


@pytest.mark.p2
def test_qa_docx_escapes_cell_html(chunkers):
    _qai_list, tbls = chunkers.qa()("table.docx", _docx_with_table())

    assert tbls
    _assert_escaped(tbls[0][0][1])


@pytest.mark.p2
def test_laws_docx_still_escapes_cell_html(chunkers):
    """Guard: the chunker that already escaped must not change."""
    chunks = chunkers.laws()("table.docx", _docx_with_table())

    table_chunk = next(chunk for chunk in chunks if "<table>" in chunk)
    _assert_escaped(table_chunk)


@pytest.mark.p2
def test_a_plain_table_is_unchanged(chunkers):
    """Guard: a table with nothing to escape renders exactly as before."""
    from docx import Document

    doc = Document()
    doc.add_heading("Heading", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Fee"
    buf = BytesIO()
    doc.save(buf)

    lines = chunkers.naive()("table.docx", buf.getvalue())
    tables = [table for _text, _image, table in lines if table]

    assert "<td>Item</td>" in tables[0]
    assert "<td>Fee</td>" in tables[0]
