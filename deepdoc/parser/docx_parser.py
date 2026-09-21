#
#  Copyright 2025 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#

import logging
import re
from collections import Counter
from io import BytesIO

import pandas as pd
from docx import Document
from docx.image.exceptions import (
    InvalidImageStreamError,
    UnexpectedEndOfFileError,
    UnrecognizedImageError,
)
from docx.oxml.ns import qn
from docx.text.run import Run

from common.constants import MAXIMUM_PAGE_NUMBER
from rag.nlp import rag_tokenizer
from rag.utils.lazy_image import LazyImage

# Markup Compatibility namespace. Word stores every text box twice inside an
# `mc:AlternateContent` element: once as a DrawingML shape under `mc:Choice` and
# once as a legacy VML shape under `mc:Fallback`. Both carry a `w:txbxContent`,
# so the fallback copy has to be skipped or the text is emitted twice.
MC_NAMESPACE = "{http://schemas.openxmlformats.org/markup-compatibility/2006}"
MC_CHOICE_TAG = MC_NAMESPACE + "Choice"
MC_FALLBACK_TAG = MC_NAMESPACE + "Fallback"
TEXT_BOX_TAG = qn("w:txbxContent")
RUN_TAG = qn("w:r")


def _has_ancestor(node, tag, stop=None):
    """Whether `node` has an ancestor with `tag`, looking no further than `stop`."""
    parent = node.getparent()
    while parent is not None and parent is not stop:
        if parent.tag == tag:
            return True
        parent = parent.getparent()
    return False


def _is_fallback_copy(box):
    """Whether `box` is the legacy copy of a text box its `mc:Choice` already holds.

    `mc:Fallback` is a general container, so a text box found only there is the
    one copy the document has and is kept.
    """
    node = box.getparent()
    while node is not None:
        if node.tag == MC_FALLBACK_TAG:
            alternate = node.getparent()
            if alternate is None:
                return False
            return any(choice.tag == MC_CHOICE_TAG and choice.find(".//" + TEXT_BOX_TAG) is not None for choice in alternate)
        node = node.getparent()
    return False


class RAGFlowDocxParser:
    @staticmethod
    def extract_text_boxes(paragraph):
        """Text of every text box anchored in `paragraph`, in document order.

        A text box keeps its own `w:p` elements in a `w:txbxContent` nested inside
        the drawing of a run, so neither `Paragraph.text` nor `Run.text` reaches it
        and callouts, pull quotes and sidebars are dropped silently.
        """
        texts = []
        element = getattr(paragraph, "_element", None)
        if element is None:
            return texts
        for child in element:
            texts.extend(RAGFlowDocxParser.text_boxes_in(child))
        return texts

    @staticmethod
    def text_boxes_in(element):
        """Text of every text box inside one child element of a paragraph."""
        texts = []
        try:
            boxes = element.findall(".//" + TEXT_BOX_TAG)
        except AttributeError:  # a test double may hand us a plain list
            return texts
        for box in boxes:
            if _is_fallback_copy(box):
                continue
            lines = []
            # Descendants, not direct children: a text box can hold a table, and
            # the text of its cells lives in `w:p` elements below the `w:tbl`.
            for p in box.iter(qn("w:p")):
                # A text box may itself contain a text box; its paragraphs are
                # collected when that inner box comes up in `boxes`.
                if _has_ancestor(p, TEXT_BOX_TAG, stop=box):
                    continue
                line = "".join(t.text or "" for t in p.iter(qn("w:t")) if not _has_ancestor(t, TEXT_BOX_TAG, stop=p))
                if line.strip():
                    lines.append(line)
            if lines:
                texts.append("\n".join(lines))
        return texts

    def get_picture(self, document, paragraph):
        imgs = paragraph._element.xpath(".//pic:pic")
        if not imgs:
            return None
        image_blobs = []
        for img in imgs:
            embed = img.xpath(".//a:blip/@r:embed")
            if not embed:
                continue
            embed = embed[0]
            image_blob = None
            try:
                related_part = document.part.related_parts[embed]
            except Exception as e:
                logging.warning(f"Skipping image due to unexpected error getting related_part: {e}")
                continue

            try:
                image = related_part.image
                if image is not None:
                    image_blob = image.blob
            except (
                UnrecognizedImageError,
                UnexpectedEndOfFileError,
                InvalidImageStreamError,
                UnicodeDecodeError,
            ) as e:
                logging.info(f"Damaged image encountered, attempting blob fallback: {e}")
            except Exception as e:
                logging.warning(f"Unexpected error getting image, attempting blob fallback: {e}")

            if image_blob is None:
                image_blob = getattr(related_part, "blob", None)
            if image_blob:
                image_blobs.append(image_blob)
        if not image_blobs:
            return None
        return LazyImage(image_blobs)

    def __extract_table_content(self, tb):
        df = []
        for row in tb.rows:
            df.append([c.text for c in row.cells])
        return self.__compose_table_content(pd.DataFrame(df))

    def __compose_table_content(self, df):

        def blockType(b):
            pattern = [
                ("^(20|19)[0-9]{2}[年/-][0-9]{1,2}[月/-][0-9]{1,2}日*$", "Dt"),
                (r"^(20|19)[0-9]{2}年$", "Dt"),
                (r"^(20|19)[0-9]{2}[年/-][0-9]{1,2}月*$", "Dt"),
                ("^[0-9]{1,2}[月/-][0-9]{1,2}日*$", "Dt"),
                (r"^第*[一二三四1-4]季度$", "Dt"),
                (r"^(20|19)[0-9]{2}年*[一二三四1-4]季度$", "Dt"),
                (r"^(20|19)[0-9]{2}[ABCDE]$", "DT"),
                ("^[0-9.,+%/ -]+$", "Nu"),
                (r"^[0-9A-Z/\._~-]+$", "Ca"),
                (r"^[A-Z]*[a-z' -]+$", "En"),
                (r"^[0-9.,+-]+[0-9A-Za-z/$￥%<>（）()' -]+$", "NE"),
                (r"^.{1}$", "Sg"),
            ]
            for p, n in pattern:
                if re.search(p, b):
                    return n
            tks = [t for t in rag_tokenizer.tokenize(b).split() if len(t) > 1]
            if len(tks) > 3:
                if len(tks) < 12:
                    return "Tx"
                else:
                    return "Lx"

            if len(tks) == 1 and rag_tokenizer.tag(tks[0]) == "nr":
                return "Nr"

            return "Ot"

        if len(df) < 2:
            return []
        max_type = Counter([blockType(str(df.iloc[i, j])) for i in range(1, len(df)) for j in range(len(df.iloc[i, :]))])
        max_type = max(max_type.items(), key=lambda x: x[1])[0]

        colnm = len(df.iloc[0, :])
        hdrows = [0]  # header is not necessarily appear in the first line
        if max_type == "Nu":
            for r in range(1, len(df)):
                tys = Counter([blockType(str(df.iloc[r, j])) for j in range(len(df.iloc[r, :]))])
                tys = max(tys.items(), key=lambda x: x[1])[0]
                if tys != max_type:
                    hdrows.append(r)

        lines = []
        for i in range(1, len(df)):
            if i in hdrows:
                continue
            hr = [r - i for r in hdrows]
            hr = [r for r in hr if r < 0]
            t = len(hr) - 1
            while t > 0:
                if hr[t] - hr[t - 1] > 1:
                    hr = hr[t:]
                    break
                t -= 1
            headers = []
            for j in range(len(df.iloc[i, :])):
                t = []
                for h in hr:
                    x = str(df.iloc[i + h, j]).strip()
                    if x in t:
                        continue
                    t.append(x)
                t = ",".join(t)
                if t:
                    t += ": "
                headers.append(t)
            cells = []
            for j in range(len(df.iloc[i, :])):
                if not str(df.iloc[i, j]):
                    continue
                cells.append(headers[j] + str(df.iloc[i, j]))
            lines.append(";".join(cells))

        if colnm > 3:
            return lines
        return ["\n".join(lines)]

    def __call__(self, fnm, from_page=0, to_page=MAXIMUM_PAGE_NUMBER):
        self.doc = Document(fnm) if isinstance(fnm, str) else Document(BytesIO(fnm))
        pn = 0  # parsed page
        secs = []  # parsed contents
        for p in self.doc.paragraphs:
            if pn > to_page:
                break

            runs_within_single_paragraph = []  # save runs within the range of pages
            boxes_within_range = []  # text boxes anchored on a page within the range
            # The paragraph's children rather than `p.runs`, so a text box is placed on
            # the page of the element it is anchored in; runs keep their old handling.
            for child in p._element:
                if pn > to_page:
                    break
                if from_page <= pn < to_page:
                    boxes_within_range.extend(self.text_boxes_in(child))
                if child.tag != RUN_TAG:
                    continue
                if from_page <= pn < to_page and p.text.strip():
                    runs_within_single_paragraph.append(Run(child, p).text)  # append run.text first

                # wrap page break checker into a static method
                if "lastRenderedPageBreak" in child.xml:
                    pn += 1

            secs.append(("".join(runs_within_single_paragraph), p.style.name if hasattr(p.style, "name") else ""))  # then concat run.text as part of the paragraph
            secs.extend((box_text, "") for box_text in boxes_within_range)

        tbls = [self.__extract_table_content(tb) for tb in self.doc.tables]
        return secs, tbls
