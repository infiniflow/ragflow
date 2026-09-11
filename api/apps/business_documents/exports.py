#
#  Copyright 2026 The InfiniFlow Authors. All Rights Reserved.
#
#  Licensed under the Apache License, Version 2.0 (the "License");
#  you may not use this file except in compliance with the License.
#  You may obtain a copy of the License at
#
#      http://www.apache.org/licenses/LICENSE-2.0
#
#  Unless required by applicable law or agreed to in writing, software
#  distributed under the License is distributed on an "AS IS" BASIS,
#  WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
#  See the License for the specific language governing permissions and
#  limitations under the License.
#

from __future__ import annotations

import hashlib
import html
import io
import logging
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Any, Callable
from urllib.parse import urlsplit

from docx import Document as WordDocument
from peewee import IntegrityError

from api.apps.business_documents.assets import published_template, rendering_policy
from api.apps.business_documents.errors import BusinessDocumentError, ConflictError, ValidationError
from api.db.db_models import BusinessDocument, BusinessDocumentExportArtifact, BusinessDocumentExportStage, BusinessDocumentJob, BusinessDocumentRevision
from business_documents.domain.workflow import OperationState
from common.misc_utils import get_uuid
from common.time_utils import current_timestamp


_FORMAT_META = {
    "MARKDOWN": ("md", "text/markdown; charset=utf-8"),
    "DOCX": ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "EVA_WIKI": ("html", "text/html; charset=utf-8"),
}

_STAGE_RESERVED = "RESERVED"
_STAGE_STORED = "STORED"
_STAGE_COMMITTED = "COMMITTED"
_STAGE_DELETE_PENDING = "DELETE_PENDING"
_STAGE_CLEANING = "CLEANING"
_CLEANUP_RETRY_MS = 60_000


def _hash_bytes(content: bytes) -> str:
    return f"sha256:{hashlib.sha256(content).hexdigest()}"


def _storage_identity(bucket: str, key: str) -> str:
    return _hash_bytes(f"{bucket}\0{key}".encode("utf-8"))


def _safe_filename(value: str) -> str:
    normalized = re.sub(r"[^\w.-]+", "_", value.strip(), flags=re.UNICODE).strip("._")
    return normalized[:160] or "business_requirements"


def _safe_external_url(value: str) -> str:
    parsed = urlsplit(value)
    if parsed.scheme.lower() not in {"http", "https"} or not parsed.netloc:
        raise ValidationError("UNSAFE_EXPORT_URL", "Only absolute HTTP(S) links are allowed in exported documents")
    return value


def _artifact_dict(row: BusinessDocumentExportArtifact) -> dict[str, Any]:
    revision = BusinessDocumentRevision.get_or_none(BusinessDocumentRevision.id == row.revision_id)
    return {
        "artifact_id": row.id,
        "document_id": row.document_id,
        "revision_id": row.revision_id,
        "revision_number": revision.revision_number if revision is not None else None,
        "format": row.export_format,
        "filename": row.filename,
        "mime_type": row.mime_type,
        "size": row.size,
        "content_hash": row.content_hash,
        "create_time": row.create_time,
    }


@dataclass(frozen=True, slots=True)
class PreparedExport:
    artifact_id: str
    document_id: str
    tenant_id: str
    owner_id: str
    revision_id: str
    revision_number: int
    export_format: str
    filename: str
    mime_type: str
    size: int
    content_hash: str
    storage_bucket: str
    storage_key: str
    create_time: int
    create_date: datetime
    created_blob: bool
    stage_id: str | None = None
    lease_token: str | None = None
    replaced_artifact_id: str | None = None
    replaced_storage_bucket: str | None = None
    replaced_storage_key: str | None = None
    replaced_content_hash: str | None = None


class BusinessDocumentExportService:
    @staticmethod
    def _for_update(query, database):
        return query.for_update() if getattr(database, "for_update", False) else query

    @classmethod
    def _lock_current_job(cls, job: BusinessDocumentJob) -> tuple[BusinessDocumentJob, BusinessDocument]:
        """Lock aggregate root then job and verify the worker attempt is current."""

        database = BusinessDocumentJob._meta.database
        document = cls._for_update(
            BusinessDocument.select().where((BusinessDocument.id == job.document_id) & (BusinessDocument.tenant_id == job.tenant_id)),
            database,
        ).first()
        current = cls._for_update(
            BusinessDocumentJob.select().where((BusinessDocumentJob.id == job.id) & (BusinessDocumentJob.document_id == job.document_id) & (BusinessDocumentJob.tenant_id == job.tenant_id)),
            database,
        ).first()
        now_ms = current_timestamp()
        if (
            document is None
            or current is None
            or current.status != "RUNNING"
            or not job.lease_token
            or current.lease_owner != job.lease_owner
            or current.lease_token != job.lease_token
            or current.lease_expires_at is None
            or current.lease_expires_at <= now_ms
        ):
            raise ConflictError("JOB_LEASE_LOST", "Export job is not held by a current worker lease")
        return current, document

    @staticmethod
    def _stage_matches(stage: BusinessDocumentExportStage, expected: dict[str, Any]) -> bool:
        return all(getattr(stage, field) == value for field, value in expected.items())

    @classmethod
    def _reserve_stage(
        cls,
        *,
        job: BusinessDocumentJob,
        document: BusinessDocument,
        revision: BusinessDocumentRevision,
        artifact_id: str,
        export_format: str,
        filename: str,
        mime_type: str,
        size: int,
        content_hash: str,
        storage_bucket: str,
        storage_key: str,
        existing: BusinessDocumentExportArtifact | None,
    ) -> BusinessDocumentExportStage:
        """Persist the exact blob intent before any object-store write."""

        database = BusinessDocumentExportStage._meta.database
        expected = {
            "job_id": job.id,
            "document_id": document.id,
            "tenant_id": document.tenant_id,
            "owner_id": document.owner_id,
            "artifact_id": artifact_id,
            "lease_token": job.lease_token,
            "revision_id": revision.id,
            "export_format": export_format,
            "filename": filename,
            "mime_type": mime_type,
            "size": size,
            "content_hash": content_hash,
            "storage_identity": _storage_identity(storage_bucket, storage_key),
            "storage_bucket": storage_bucket,
            "storage_key": storage_key,
            "replaced_artifact_id": existing.id if existing is not None else None,
            "replaced_storage_bucket": existing.storage_bucket if existing is not None else None,
            "replaced_storage_key": existing.storage_key if existing is not None else None,
            "replaced_content_hash": existing.content_hash if existing is not None else None,
        }
        with database.atomic():
            current_job, current_document = cls._lock_current_job(job)
            command_payload = current_job.payload.get("command_payload", {})
            if (
                current_document.lifecycle_state != "AGREED"
                or current_document.operation_state != OperationState.EXPORTING.value
                or current_document.current_revision_id != revision.id
                or current_document.state_version != current_job.source_state_version
                or command_payload.get("revision_id") != revision.id
                or command_payload.get("format") != export_format
                or current_document.owner_id != document.owner_id
            ):
                raise ConflictError("EXPORT_JOB_CHANGED", "Export job changed before its blob could be staged")
            query = BusinessDocumentExportStage.select().where((BusinessDocumentExportStage.job_id == job.id) & (BusinessDocumentExportStage.lease_token == job.lease_token))
            stage = cls._for_update(query, database).first()
            if stage is None:
                timestamp = current_timestamp()
                now = datetime.now()
                try:
                    with database.atomic():
                        stage = BusinessDocumentExportStage.create(
                            id=get_uuid(),
                            state=_STAGE_RESERVED,
                            create_time=timestamp,
                            create_date=now,
                            update_time=timestamp,
                            update_date=now,
                            **expected,
                        )
                except IntegrityError:
                    stage = cls._for_update(
                        BusinessDocumentExportStage.select().where((BusinessDocumentExportStage.job_id == job.id) & (BusinessDocumentExportStage.lease_token == job.lease_token)),
                        database,
                    ).first()
            if stage is None or stage.state not in {_STAGE_RESERVED, _STAGE_STORED} or not cls._stage_matches(stage, expected):
                raise ConflictError("EXPORT_STAGE_CHANGED", "The durable export stage does not match the current worker attempt")
            return stage

    @classmethod
    def _mark_stage_stored(cls, stage_id: str, job: BusinessDocumentJob) -> BusinessDocumentExportStage:
        """Fence the stage after the object bytes have been read back and hashed."""

        database = BusinessDocumentExportStage._meta.database
        with database.atomic():
            cls._lock_current_job(job)
            stage = cls._for_update(
                BusinessDocumentExportStage.select().where((BusinessDocumentExportStage.id == stage_id) & (BusinessDocumentExportStage.lease_token == job.lease_token)),
                database,
            ).first()
            if stage is None or stage.state not in {_STAGE_RESERVED, _STAGE_STORED}:
                raise ConflictError("EXPORT_STAGE_CHANGED", "The durable export stage is no longer writable")
            timestamp = current_timestamp()
            changed = (
                BusinessDocumentExportStage.update(
                    state=_STAGE_STORED,
                    update_time=timestamp,
                    update_date=datetime.now(),
                )
                .where(
                    (BusinessDocumentExportStage.id == stage.id)
                    & (BusinessDocumentExportStage.lease_token == job.lease_token)
                    & (BusinessDocumentExportStage.state.in_((_STAGE_RESERVED, _STAGE_STORED)))
                )
                .execute()
            )
            if changed != 1:
                raise ConflictError("EXPORT_STAGE_CHANGED", "The durable export stage changed while bytes were stored")
            stage.state = _STAGE_STORED
            return stage

    @staticmethod
    def _blob_referenced(bucket: str, key: str) -> bool:
        return BusinessDocumentExportArtifact.select().where((BusinessDocumentExportArtifact.storage_bucket == bucket) & (BusinessDocumentExportArtifact.storage_key == key)).exists()

    @staticmethod
    def _remove_blob_verified(storage, bucket: str, key: str) -> bool:
        """Remove a blob and retain its ledger whenever absence cannot be trusted."""

        try:
            remove_and_confirm_absent = getattr(storage, "remove_and_confirm_absent", None)
            if not callable(remove_and_confirm_absent):
                raise RuntimeError("Storage does not provide strict object-removal verification")
            return remove_and_confirm_absent(bucket, key) is True
        except Exception:
            logging.exception("Unable to verify removal of business document export blob %s/%s", bucket, key)
            return False

    @staticmethod
    def _reschedule_stage(stage: BusinessDocumentExportStage, cleanup_after: int) -> None:
        """Delay an inconclusive row without losing its cleanup evidence."""

        try:
            BusinessDocumentExportStage.update(
                cleanup_after=cleanup_after,
            ).where((BusinessDocumentExportStage.id == stage.id) & (BusinessDocumentExportStage.storage_identity == stage.storage_identity)).execute()
        except Exception:
            logging.exception("Unable to reschedule business document export stage %s", stage.id)

    @classmethod
    def _claim_stage_cleanup(
        cls,
        stage_id: str,
        *,
        now_ms: int,
        defer_live: bool,
    ) -> tuple[BusinessDocumentExportStage | None, bool]:
        """Fence publication before doing an idempotent cross-store cleanup."""

        snapshot = BusinessDocumentExportStage.get_or_none(BusinessDocumentExportStage.id == stage_id)
        if snapshot is None:
            return None, False
        database = BusinessDocumentExportStage._meta.database
        with database.atomic():
            cls._for_update(
                BusinessDocument.select().where(BusinessDocument.id == snapshot.document_id),
                database,
            ).first()
            job = None
            if snapshot.job_id:
                job = cls._for_update(
                    BusinessDocumentJob.select().where(BusinessDocumentJob.id == snapshot.job_id),
                    database,
                ).first()
            stage = cls._for_update(
                BusinessDocumentExportStage.select().where(BusinessDocumentExportStage.id == stage_id),
                database,
            ).first()
            if stage is None:
                return None, False
            if defer_live and (
                stage.state in {_STAGE_RESERVED, _STAGE_STORED}
                and job is not None
                and job.status == "RUNNING"
                and job.lease_token == stage.lease_token
                and job.lease_expires_at is not None
                and job.lease_expires_at > now_ms
            ):
                cls._reschedule_stage(stage, job.lease_expires_at)
                return None, True
            if stage.state != _STAGE_CLEANING:
                timestamp = current_timestamp()
                changed = (
                    BusinessDocumentExportStage.update(
                        state=_STAGE_CLEANING,
                        update_time=timestamp,
                        update_date=datetime.now(),
                    )
                    .where((BusinessDocumentExportStage.id == stage.id) & (BusinessDocumentExportStage.state == stage.state) & (BusinessDocumentExportStage.storage_identity == stage.storage_identity))
                    .execute()
                )
                if changed != 1:
                    return None, False
                stage.state = _STAGE_CLEANING
            return stage, False

    @classmethod
    def _cleanup_claimed_stage(cls, stage: BusinessDocumentExportStage, storage) -> bool:
        candidates = [(stage.storage_bucket, stage.storage_key)]
        if stage.replaced_storage_bucket and stage.replaced_storage_key:
            candidates.append((stage.replaced_storage_bucket, stage.replaced_storage_key))
        candidates = list(dict.fromkeys(candidates))
        removed: set[tuple[str, str]] = set()
        for bucket, key in candidates:
            try:
                referenced = cls._blob_referenced(bucket, key)
            except Exception:
                logging.exception("Unable to verify business document export blob reference %s/%s", bucket, key)
                return False
            if referenced:
                continue
            if storage is None or not cls._remove_blob_verified(storage, bucket, key):
                return False
            removed.add((bucket, key))
        database = BusinessDocumentExportStage._meta.database
        with database.atomic():
            current = cls._for_update(
                BusinessDocumentExportStage.select().where(BusinessDocumentExportStage.id == stage.id),
                database,
            ).first()
            if current is None:
                return True
            if current.state != _STAGE_CLEANING or current.storage_identity != stage.storage_identity:
                return False
            for bucket, key in candidates:
                if (bucket, key) not in removed and not cls._blob_referenced(bucket, key):
                    return False
            return (
                BusinessDocumentExportStage.delete()
                .where(
                    (BusinessDocumentExportStage.id == current.id) & (BusinessDocumentExportStage.state == _STAGE_CLEANING) & (BusinessDocumentExportStage.storage_identity == current.storage_identity)
                )
                .execute()
                == 1
            )

    @classmethod
    def _cleanup_stage_id(cls, stage_id: str, storage, *, now_ms: int, defer_live: bool) -> str:
        stage, deferred = cls._claim_stage_cleanup(stage_id, now_ms=now_ms, defer_live=defer_live)
        if deferred:
            return "deferred"
        if stage is None:
            current = BusinessDocumentExportStage.get_or_none(BusinessDocumentExportStage.id == stage_id)
            if current is None:
                return "cleaned"
            cls._reschedule_stage(current, now_ms + _CLEANUP_RETRY_MS)
            return "failures"
        if cls._cleanup_claimed_stage(stage, storage):
            return "cleaned"
        cls._reschedule_stage(stage, now_ms + _CLEANUP_RETRY_MS)
        return "failures"

    @classmethod
    def cleanup_stage(cls, stage_id: str, storage=None) -> bool:
        storage_impl = storage
        if storage_impl is None:
            try:
                storage_impl = cls._default_storage()
            except Exception:
                logging.exception("Business document export storage is unavailable during stage cleanup")
        return cls._cleanup_stage_id(stage_id, storage_impl, now_ms=current_timestamp(), defer_live=False) == "cleaned"

    @classmethod
    def reconcile_staging(
        cls,
        storage=None,
        *,
        document_id: str | None = None,
        limit: int = 100,
        now_ms: int | None = None,
    ) -> dict[str, int]:
        """Reconcile abandoned/committed stages without touching a live attempt."""

        storage_impl = storage
        if storage_impl is None:
            try:
                storage_impl = cls._default_storage()
            except Exception:
                logging.exception("Business document export storage is unavailable during reconciliation")
        now_ms = current_timestamp() if now_ms is None else now_ms
        query = (
            BusinessDocumentExportStage.select()
            .where(BusinessDocumentExportStage.cleanup_after <= now_ms)
            .order_by(BusinessDocumentExportStage.cleanup_after.asc(), BusinessDocumentExportStage.create_time.asc(), BusinessDocumentExportStage.id.asc())
            .limit(max(1, limit))
        )
        if document_id is not None:
            query = query.where(BusinessDocumentExportStage.document_id == document_id)
        result = {"scanned": 0, "cleaned": 0, "deferred": 0, "failures": 0}
        for stage in list(query):
            result["scanned"] += 1
            outcome = cls._cleanup_stage_id(stage.id, storage_impl, now_ms=now_ms, defer_live=True)
            result[outcome] += 1
        return result

    @classmethod
    def queue_artifact_cleanup(cls, artifact: BusinessDocumentExportArtifact) -> BusinessDocumentExportStage:
        """Durably record artifact deletion before its metadata is removed."""

        database = BusinessDocumentExportStage._meta.database
        identity = _storage_identity(artifact.storage_bucket, artifact.storage_key)
        with database.atomic():
            stage = cls._for_update(
                BusinessDocumentExportStage.select().where(BusinessDocumentExportStage.storage_identity == identity),
                database,
            ).first()
            timestamp = current_timestamp()
            now = datetime.now()
            if stage is None:
                stage = BusinessDocumentExportStage.create(
                    id=get_uuid(),
                    job_id=None,
                    document_id=artifact.document_id,
                    tenant_id=artifact.tenant_id,
                    owner_id=artifact.owner_id,
                    artifact_id=artifact.id,
                    lease_token=None,
                    revision_id=artifact.revision_id,
                    export_format=artifact.export_format,
                    filename=artifact.filename,
                    mime_type=artifact.mime_type,
                    size=artifact.size,
                    content_hash=artifact.content_hash,
                    storage_identity=identity,
                    storage_bucket=artifact.storage_bucket,
                    storage_key=artifact.storage_key,
                    state=_STAGE_DELETE_PENDING,
                    cleanup_after=0,
                    create_time=timestamp,
                    create_date=now,
                    update_time=timestamp,
                    update_date=now,
                )
            else:
                expected = {
                    "document_id": artifact.document_id,
                    "tenant_id": artifact.tenant_id,
                    "owner_id": artifact.owner_id,
                    "artifact_id": artifact.id,
                    "revision_id": artifact.revision_id,
                    "export_format": artifact.export_format,
                    "filename": artifact.filename,
                    "mime_type": artifact.mime_type,
                    "size": artifact.size,
                    "content_hash": artifact.content_hash,
                    "storage_bucket": artifact.storage_bucket,
                    "storage_key": artifact.storage_key,
                }
                if not cls._stage_matches(stage, expected):
                    raise ConflictError("EXPORT_STAGE_CHANGED", "The cleanup ledger does not match the stored artifact")
                BusinessDocumentExportStage.update(
                    state=_STAGE_DELETE_PENDING,
                    cleanup_after=0,
                    update_time=timestamp,
                    update_date=now,
                ).where(BusinessDocumentExportStage.id == stage.id).execute()
                stage.state = _STAGE_DELETE_PENDING
            return stage

    @classmethod
    def generate(
        cls,
        job: BusinessDocumentJob,
        storage=None,
        *,
        ensure_current: Callable[[], None] | None = None,
    ) -> PreparedExport:
        """Render and durably stage bytes without publishing artifact metadata."""

        if job.job_type != "GENERATE_EXPORT":
            raise ValidationError("INVALID_EXPORT_JOB", "Job is not an export request")
        if job.status != "RUNNING" or not job.lease_token:
            raise ConflictError("JOB_LEASE_LOST", "Export job is not held by a current worker lease")
        document = BusinessDocument.get_or_none((BusinessDocument.id == job.document_id) & (BusinessDocument.tenant_id == job.tenant_id))
        if document is None:
            raise BusinessDocumentError("DOCUMENT_NOT_FOUND", "Business document not found", 404)
        command_payload = job.payload.get("command_payload", {})
        revision_id = command_payload.get("revision_id")
        export_format = command_payload.get("format")
        if document.lifecycle_state != "AGREED" or document.operation_state != OperationState.EXPORTING.value or document.current_revision_id != revision_id:
            raise ConflictError("AGREED_REVISION_REQUIRED", "Export job no longer targets the current agreed revision")
        if document.state_version != job.source_state_version:
            raise ConflictError("STALE_AI_RESULT", "Export job targets an outdated document state")
        if export_format not in _FORMAT_META:
            raise ValidationError("INVALID_EXPORT_FORMAT", "Only MARKDOWN, DOCX and EVA_WIKI are supported")
        revision = BusinessDocumentRevision.get_or_none((BusinessDocumentRevision.id == revision_id) & (BusinessDocumentRevision.document_id == document.id))
        if revision is None:
            raise BusinessDocumentError("REVISION_NOT_FOUND", "Business document revision not found", 404)

        storage_impl = storage or cls._default_storage()
        if ensure_current is not None:
            ensure_current()
        existing = BusinessDocumentExportArtifact.get_or_none(
            (BusinessDocumentExportArtifact.document_id == document.id) & (BusinessDocumentExportArtifact.revision_id == revision.id) & (BusinessDocumentExportArtifact.export_format == export_format)
        )
        if existing is not None:
            try:
                existing_content = storage_impl.get(existing.storage_bucket, existing.storage_key)
            except Exception:
                existing_content = None
            if isinstance(existing_content, bytes) and _hash_bytes(existing_content) == existing.content_hash:
                if ensure_current is not None:
                    ensure_current()
                return PreparedExport(
                    artifact_id=existing.id,
                    document_id=existing.document_id,
                    tenant_id=existing.tenant_id,
                    owner_id=existing.owner_id,
                    revision_id=existing.revision_id,
                    revision_number=revision.revision_number,
                    export_format=existing.export_format,
                    filename=existing.filename,
                    mime_type=existing.mime_type,
                    size=existing.size,
                    content_hash=existing.content_hash,
                    storage_bucket=existing.storage_bucket,
                    storage_key=existing.storage_key,
                    create_time=existing.create_time,
                    create_date=existing.create_date,
                    created_blob=False,
                    lease_token=job.lease_token,
                )

        content = cls._render(export_format, revision)
        if ensure_current is not None:
            ensure_current()
        extension, mime_type = _FORMAT_META[export_format]
        artifact_id = job.id
        filename = f"{_safe_filename(document.title)}_r{revision.revision_number}.{extension}"
        bucket = f"{document.tenant_id}-business-documents"
        storage_key = f"exports/{document.id}/{revision.id}/{artifact_id}/{job.lease_token}.{extension}"
        content_hash = _hash_bytes(content)
        stage = cls._reserve_stage(
            job=job,
            document=document,
            revision=revision,
            artifact_id=artifact_id,
            export_format=export_format,
            filename=filename,
            mime_type=mime_type,
            size=len(content),
            content_hash=content_hash,
            storage_bucket=bucket,
            storage_key=storage_key,
            existing=existing,
        )
        stored_content = None
        if stage.state == _STAGE_STORED:
            try:
                stored_content = storage_impl.get(bucket, storage_key)
            except Exception:
                stored_content = None
        if not isinstance(stored_content, bytes) or _hash_bytes(stored_content) != content_hash:
            try:
                storage_impl.put(bucket, storage_key, content)
                stored_content = storage_impl.get(bucket, storage_key)
            except Exception as exc:
                raise BusinessDocumentError("EXPORT_STORAGE_WRITE_FAILED", "Export artifact could not be durably stored", 500) from exc
        if not isinstance(stored_content, bytes) or _hash_bytes(stored_content) != content_hash:
            raise BusinessDocumentError("EXPORT_STORAGE_WRITE_FAILED", "Export artifact could not be durably stored", 500)
        if ensure_current is not None:
            ensure_current()
        stage = cls._mark_stage_stored(stage.id, job)
        return PreparedExport(
            artifact_id=artifact_id,
            document_id=document.id,
            tenant_id=document.tenant_id,
            owner_id=document.owner_id,
            revision_id=revision.id,
            revision_number=revision.revision_number,
            export_format=export_format,
            filename=filename,
            mime_type=mime_type,
            size=len(content),
            content_hash=content_hash,
            storage_bucket=bucket,
            storage_key=storage_key,
            create_time=stage.create_time,
            create_date=stage.create_date,
            created_blob=True,
            stage_id=stage.id,
            lease_token=job.lease_token,
            replaced_artifact_id=existing.id if existing is not None else None,
            replaced_storage_bucket=existing.storage_bucket if existing is not None else None,
            replaced_storage_key=existing.storage_key if existing is not None else None,
            replaced_content_hash=existing.content_hash if existing is not None else None,
        )

    @classmethod
    def commit(
        cls,
        prepared: object,
        *,
        document: BusinessDocument,
        job: BusinessDocumentJob,
    ) -> dict[str, Any]:
        """Publish prepared metadata inside the caller's fenced DB transaction."""

        if not isinstance(prepared, PreparedExport):
            raise ValidationError("INVALID_EXPORT_RESULT", "Prepared export is required")
        command_payload = job.payload.get("command_payload", {})
        format_meta = _FORMAT_META.get(prepared.export_format)
        expected_extension = format_meta[0] if format_meta is not None else None
        expected_mime_type = format_meta[1] if format_meta is not None else None
        expected_bucket = f"{document.tenant_id}-business-documents"
        expected_storage_key = f"exports/{document.id}/{prepared.revision_id}/{prepared.artifact_id}/{prepared.lease_token}.{expected_extension}"
        if (
            format_meta is None
            or job.job_type != "GENERATE_EXPORT"
            or document.lifecycle_state != "AGREED"
            or document.operation_state != OperationState.EXPORTING.value
            or document.state_version != job.source_state_version
            or prepared.document_id != document.id
            or prepared.tenant_id != document.tenant_id
            or (prepared.created_blob and prepared.owner_id != document.owner_id)
            or prepared.revision_id != document.current_revision_id
            or prepared.revision_id != command_payload.get("revision_id")
            or prepared.export_format != command_payload.get("format")
            or prepared.lease_token != job.lease_token
            or (prepared.created_blob and prepared.artifact_id != job.id)
            or (prepared.created_blob and prepared.storage_bucket != expected_bucket)
            or (prepared.created_blob and prepared.storage_key != expected_storage_key)
            or (prepared.created_blob and prepared.mime_type != expected_mime_type)
            or not prepared.content_hash.startswith("sha256:")
            or len(prepared.content_hash) != 71
        ):
            raise ConflictError("EXPORT_JOB_CHANGED", "Prepared export no longer matches the current export job")
        revision = BusinessDocumentRevision.get_or_none((BusinessDocumentRevision.id == prepared.revision_id) & (BusinessDocumentRevision.document_id == document.id))
        if revision is None or revision.revision_number != prepared.revision_number:
            raise BusinessDocumentError("REVISION_NOT_FOUND", "Business document revision not found", 404)

        current = BusinessDocumentExportArtifact.get_or_none(
            (BusinessDocumentExportArtifact.document_id == document.id)
            & (BusinessDocumentExportArtifact.revision_id == prepared.revision_id)
            & (BusinessDocumentExportArtifact.export_format == prepared.export_format)
        )
        if not prepared.created_blob:
            if prepared.stage_id is not None:
                raise ConflictError("EXPORT_STAGE_CHANGED", "A reused artifact must not carry a mutable staging record")
            if (
                current is None
                or current.id != prepared.artifact_id
                or current.storage_bucket != prepared.storage_bucket
                or current.storage_key != prepared.storage_key
                or current.content_hash != prepared.content_hash
            ):
                raise ConflictError("EXPORT_ARTIFACT_CHANGED", "Export artifact changed before completion")
            return _artifact_dict(current)

        if prepared.stage_id is None:
            raise ConflictError("EXPORT_STAGE_CHANGED", "A durable staging record is required before artifact publication")
        database = BusinessDocumentExportStage._meta.database
        stage = cls._for_update(
            BusinessDocumentExportStage.select().where(BusinessDocumentExportStage.id == prepared.stage_id),
            database,
        ).first()
        expected_stage = {
            "job_id": job.id,
            "document_id": document.id,
            "tenant_id": document.tenant_id,
            "owner_id": document.owner_id,
            "artifact_id": prepared.artifact_id,
            "lease_token": job.lease_token,
            "revision_id": prepared.revision_id,
            "export_format": prepared.export_format,
            "filename": prepared.filename,
            "mime_type": prepared.mime_type,
            "size": prepared.size,
            "content_hash": prepared.content_hash,
            "storage_identity": _storage_identity(prepared.storage_bucket, prepared.storage_key),
            "storage_bucket": prepared.storage_bucket,
            "storage_key": prepared.storage_key,
            "replaced_artifact_id": prepared.replaced_artifact_id,
            "replaced_storage_bucket": prepared.replaced_storage_bucket,
            "replaced_storage_key": prepared.replaced_storage_key,
            "replaced_content_hash": prepared.replaced_content_hash,
        }
        if stage is None or stage.state != _STAGE_STORED or not cls._stage_matches(stage, expected_stage):
            raise ConflictError("EXPORT_STAGE_CHANGED", "The durable export stage is not publishable")

        if prepared.replaced_artifact_id is None:
            if current is not None:
                raise ConflictError("EXPORT_ARTIFACT_CHANGED", "Export artifact appeared before completion")
        else:
            if (
                current is None
                or current.id != prepared.replaced_artifact_id
                or current.storage_bucket != prepared.replaced_storage_bucket
                or current.storage_key != prepared.replaced_storage_key
                or current.content_hash != prepared.replaced_content_hash
            ):
                raise ConflictError("EXPORT_ARTIFACT_CHANGED", "Export artifact changed before repair completion")
            deleted = (
                BusinessDocumentExportArtifact.delete()
                .where(
                    (BusinessDocumentExportArtifact.id == current.id)
                    & (BusinessDocumentExportArtifact.content_hash == current.content_hash)
                    & (BusinessDocumentExportArtifact.storage_key == current.storage_key)
                )
                .execute()
            )
            if deleted != 1:
                raise ConflictError("EXPORT_ARTIFACT_CHANGED", "Export artifact changed while it was being repaired")

        artifact = BusinessDocumentExportArtifact.create(
            id=prepared.artifact_id,
            document_id=document.id,
            tenant_id=document.tenant_id,
            owner_id=document.owner_id,
            revision_id=revision.id,
            export_format=prepared.export_format,
            filename=prepared.filename,
            mime_type=prepared.mime_type,
            size=prepared.size,
            content_hash=prepared.content_hash,
            storage_bucket=prepared.storage_bucket,
            storage_key=prepared.storage_key,
            create_time=prepared.create_time,
            create_date=prepared.create_date,
            update_time=prepared.create_time,
            update_date=prepared.create_date,
        )
        timestamp = current_timestamp()
        changed = (
            BusinessDocumentExportStage.update(
                state=_STAGE_COMMITTED,
                cleanup_after=0,
                update_time=timestamp,
                update_date=datetime.now(),
            )
            .where((BusinessDocumentExportStage.id == stage.id) & (BusinessDocumentExportStage.state == _STAGE_STORED) & (BusinessDocumentExportStage.lease_token == job.lease_token))
            .execute()
        )
        if changed != 1:
            raise ConflictError("EXPORT_STAGE_CHANGED", "The durable export stage changed during publication")
        return _artifact_dict(artifact)

    @classmethod
    def discard(cls, prepared: object, storage=None) -> None:
        """Remove only blobs which no committed artifact metadata references."""

        if not isinstance(prepared, PreparedExport):
            return
        if prepared.stage_id is not None:
            try:
                cls.cleanup_stage(prepared.stage_id, storage=storage)
            except Exception:
                logging.exception("Unable to reconcile business document export stage %s", prepared.stage_id)
            return
        storage_impl = storage
        if storage_impl is None:
            try:
                storage_impl = cls._default_storage()
            except Exception:
                logging.exception("Business document export storage is unavailable during fallback cleanup")
                return
        candidates: list[tuple[str, str]] = []
        if prepared.created_blob:
            candidates.append((prepared.storage_bucket, prepared.storage_key))
        if prepared.replaced_storage_bucket and prepared.replaced_storage_key:
            candidates.append((prepared.replaced_storage_bucket, prepared.replaced_storage_key))
        for bucket, key in candidates:
            try:
                referenced = BusinessDocumentExportArtifact.select().where((BusinessDocumentExportArtifact.storage_bucket == bucket) & (BusinessDocumentExportArtifact.storage_key == key)).exists()
            except Exception:
                logging.exception("Unable to verify business document export blob reference %s/%s", bucket, key)
                continue
            if referenced:
                continue
            cls._remove_blob_verified(storage_impl, bucket, key)

    @classmethod
    def list_artifacts(cls, tenant_id: str, actor_id: str, document_id: str, is_admin: bool = False) -> list[dict[str, Any]]:
        if not BusinessDocument.select().where(BusinessDocument.id == document_id).exists():
            raise BusinessDocumentError("DOCUMENT_NOT_FOUND", "Business document not found", 404)
        rows = BusinessDocumentExportArtifact.select().where(BusinessDocumentExportArtifact.document_id == document_id).order_by(BusinessDocumentExportArtifact.create_time.desc())
        return [_artifact_dict(row) for row in rows]

    @classmethod
    def download(cls, tenant_id: str, actor_id: str, document_id: str, artifact_id: str, storage=None, is_admin: bool = False):
        artifact_query = (BusinessDocumentExportArtifact.id == artifact_id) & (BusinessDocumentExportArtifact.document_id == document_id)
        if not BusinessDocument.select().where(BusinessDocument.id == document_id).exists():
            raise BusinessDocumentError("DOCUMENT_NOT_FOUND", "Business document not found", 404)
        artifact = BusinessDocumentExportArtifact.get_or_none(artifact_query)
        if artifact is None:
            raise BusinessDocumentError("EXPORT_NOT_FOUND", "Export artifact not found", 404)
        storage_impl = storage or cls._default_storage()
        try:
            content = storage_impl.get(artifact.storage_bucket, artifact.storage_key)
        except Exception as exc:
            raise BusinessDocumentError("EXPORT_STORAGE_CORRUPT", "Stored export artifact is missing or corrupt", 500) from exc
        if not isinstance(content, bytes) or _hash_bytes(content) != artifact.content_hash:
            raise BusinessDocumentError("EXPORT_STORAGE_CORRUPT", "Stored export artifact is missing or corrupt", 500)
        return _artifact_dict(artifact), content

    @staticmethod
    def _default_storage():
        from common import settings
        from api.apps.business_documents.adapters.storage import BusinessDocumentStorageAdapter

        if settings.STORAGE_IMPL is None:
            raise RuntimeError("RAGFlow storage is not initialized")
        return BusinessDocumentStorageAdapter(settings.STORAGE_IMPL, settings.STORAGE_IMPL_TYPE)

    @classmethod
    def _render(cls, export_format: str, revision: BusinessDocumentRevision) -> bytes:
        if export_format == "MARKDOWN":
            return revision.body_markdown.encode("utf-8")
        if export_format == "DOCX":
            return cls._render_docx(revision.document_ast)
        if export_format == "EVA_WIKI":
            return cls._render_eva_wiki(revision.document_ast, revision.revision_number).encode("utf-8")
        raise ValidationError("INVALID_EXPORT_FORMAT", "Unsupported export format")

    @staticmethod
    def _render_docx(document_ast: dict[str, Any]) -> bytes:
        document = WordDocument()
        for section in document_ast["sections"]:
            level = min(9, 1 + section["id"].count("."))
            document.add_heading(f"{section['id']}. {section['title']}", level=level)
            for block in section["blocks"]:
                block_type = block["type"]
                if block_type == "paragraph":
                    document.add_paragraph(block["text"])
                elif block_type == "list":
                    for item in block["items"]:
                        document.add_paragraph(str(item), style="List Bullet")
                elif block_type == "table":
                    table = document.add_table(rows=1, cols=len(block["headers"]))
                    for index, header in enumerate(block["headers"]):
                        table.rows[0].cells[index].text = str(header)
                    for row in block["rows"]:
                        cells = table.add_row().cells
                        for index, value in enumerate(row[: len(cells)]):
                            cells[index].text = "" if value is None else str(value)
                elif block_type == "plantuml":
                    document.add_paragraph(block["source"], style="No Spacing")
                elif block_type in {"image", "reference"}:
                    document.add_paragraph(f"{block.get('alt') or block.get('label')}: {block['url']}")
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _render_eva_wiki(document_ast: dict[str, Any], revision_number: int) -> str:
        policy = rendering_policy()["eva_wiki"]
        if not policy.get("exclude_protocol"):
            raise RuntimeError("EvaWiki export policy must exclude review protocol")
        generated_at = datetime.now(UTC).date().isoformat()
        lines = [
            f'<div class="business-requirements" data-template-version="{html.escape(published_template()["template_version"])}" data-revision="{revision_number}" data-generated-at="{generated_at}">'
        ]
        for section_index, section in enumerate(document_ast["sections"]):
            heading_level = min(6, 2 + section["id"].count("."))
            section_id = html.escape(section["id"])
            section_node_id = f"br-r{revision_number}-s{section_index + 1}"
            lines.append(f'<h{heading_level} data-id="{section_node_id}">{section_id}. {html.escape(section["title"])}</h{heading_level}>')
            for block_index, block in enumerate(section["blocks"]):
                block_type = block["type"]
                block_node_id = f"{section_node_id}-b{block_index + 1}"
                if block_type == "paragraph":
                    if str(block["text"]).strip():
                        lines.append(f'<p data-id="{block_node_id}">{html.escape(block["text"])}</p>')
                elif block_type == "list":
                    items = "".join(
                        f'<li data-id="{block_node_id}-li{item_index}"><p data-id="{block_node_id}-li{item_index}-p">{html.escape(str(item))}</p></li>'
                        for item_index, item in enumerate(block["items"], start=1)
                        if str(item).strip()
                    )
                    if items:
                        lines.append(f'<ul class="{policy["root_list_class"]}" style="list-style-type: disc;" data-id="{block_node_id}">{items}</ul>')
                elif block_type == "table":
                    headers = "".join(
                        f'<th colspan="1" rowspan="1" data-x="{x}" data-y="0" data-id="{block_node_id}-h{x}"><p data-id="{block_node_id}-h{x}-p">{html.escape(str(item)) or "&#160;"}</p></th>'
                        for x, item in enumerate(block["headers"])
                    )
                    rows = "".join(
                        f'<tr data-id="{block_node_id}-r{y}">'
                        + "".join(
                            f'<td colspan="1" rowspan="1" data-x="{x}" data-y="{y}" data-id="{block_node_id}-c{x}-{y}">'
                            f'<p data-id="{block_node_id}-c{x}-{y}-p">'
                            f"{html.escape('' if value is None else str(value)) or '&#160;'}</p></td>"
                            for x, value in enumerate(row)
                        )
                        + "</tr>"
                        for y, row in enumerate(block["rows"], start=1)
                    )
                    lines.append(
                        f'<div class="{policy["table_wrapper_class"]}" data-macros="{policy["table_wrapper_macro"]}" '
                        f'data-id="{block_node_id}"><table data-id="{block_node_id}-table">'
                        f'<thead><tr data-id="{block_node_id}-r0">{headers}</tr></thead><tbody>{rows}</tbody></table></div>'
                    )
                elif block_type == "plantuml":
                    lines.append(f'<pre class="{policy["plantuml_class"]}" data-id="{block_node_id}"><code data-id="{block_node_id}-code">{html.escape(block["source"])}</code></pre>')
                elif block_type == "image":
                    image_url = _safe_external_url(block["url"])
                    lines.append(f'<img data-id="{block_node_id}" alt="{html.escape(block["alt"])}" src="{html.escape(image_url)}" />')
                elif block_type == "reference":
                    reference_url = _safe_external_url(block["url"])
                    lines.append(f'<p data-id="{block_node_id}"><a data-id="{block_node_id}-a" href="{html.escape(reference_url)}">{html.escape(block["label"])}</a></p>')
        lines.append("</div>")
        return "\n".join(lines)
